"""Report assembly + export (Module 11).

Builds a structured markdown report from selected data sources, optionally
refined by the LLM. Never invents numbers: when a source has no data it says so.

Export formats: Markdown, HTML, DOCX, and PDF (reportlab; Farsi is reshaped).
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import markdown as md_lib

REPORT_SECTIONS = [
    "Executive Summary",
    "Main Findings",
    "Important Events",
    "Risks",
    "Issues",
    "Statistics",
    "Recommendations",
    "Action Items",
    "Data Sources",
]


def build_skeleton(
    title: str,
    report_type: str,
    date_from: str,
    date_to: str,
    sources: dict[str, str],
    language: str = "en",
) -> str:
    """Assemble a markdown skeleton. `sources` maps a source name -> its content
    (or an empty string, which is rendered as an explicit 'no data' note)."""
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    lines = [f"# {title}", "", f"*{report_type} · {date_from} → {date_to} · generated {now}*", ""]
    for section in REPORT_SECTIONS:
        lines.append(f"## {section}")
        lines.append("")
        if section == "Data Sources":
            for name, content in sources.items():
                status = "included" if content.strip() else "no data available"
                lines.append(f"- **{name}**: {status}")
            lines.append("")
        else:
            lines.append("_To be completed._" if language == "en" else "_تکمیل شود._")
            lines.append("")
    lines.append("## Appendix: Source Data")
    lines.append("")
    for name, content in sources.items():
        lines.append(f"### {name}")
        lines.append("")
        lines.append(content.strip() if content.strip() else "_No data available for this source._")
        lines.append("")
    return "\n".join(lines)


def build_llm_instruction(title: str, report_type: str, language: str, detail: str) -> str:
    lang = "Persian (فارسی)" if language == "fa" else "English"
    return (
        f"You are writing a {report_type} titled '{title}' in {lang}. "
        f"Detail level: {detail}. Use ONLY the provided source data. "
        "Do not invent numbers or facts. When a section has no supporting data, "
        "state that the information is unavailable. Produce clean Markdown with "
        f"these sections: {', '.join(REPORT_SECTIONS)}."
    )


def export_markdown(content: str, out_dir: Path, name: str) -> Path:
    path = out_dir / f"{name}.md"
    path.write_text(content, encoding="utf-8")
    return path


def export_html(content: str, out_dir: Path, name: str, title: str = "") -> Path:
    body = md_lib.markdown(content, extensions=["tables", "fenced_code", "toc"])
    html = f"""<!doctype html>
<html lang="en" dir="auto">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title or name}</title>
<style>
 body {{ font-family: system-ui, 'Segoe UI', Tahoma, sans-serif; max-width: 860px;
        margin: 2rem auto; padding: 0 1rem; line-height: 1.6; color: #1a1a2e; }}
 table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; }}
 th, td {{ border: 1px solid #ccc; padding: 6px 10px; text-align: start; }}
 th {{ background: #f2f4f8; }} code {{ background: #f2f4f8; padding: 2px 4px; border-radius: 3px; }}
 pre {{ background: #1a1a2e; color: #e6e6e6; padding: 1rem; overflow-x: auto; border-radius: 6px; }}
 h1,h2,h3 {{ color: #16213e; }}
</style>
</head>
<body>{body}</body>
</html>"""
    path = out_dir / f"{name}.html"
    path.write_text(html, encoding="utf-8")
    return path


def export_docx(content: str, out_dir: Path, name: str) -> Path:
    from docx import Document

    doc = Document()
    for line in content.splitlines():
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
    path = out_dir / f"{name}.docx"
    doc.save(str(path))
    return path


def export_pdf(content: str, out_dir: Path, name: str, title: str = "", language: str = "en") -> Path:
    """Render Markdown to a real PDF (reportlab). Farsi is reshaped/bidi-reordered."""
    from .pdf_export import export_markdown_pdf

    return export_markdown_pdf(content, out_dir, name, title=title, language=language)
