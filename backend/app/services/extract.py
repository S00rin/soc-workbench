"""Extract text/markdown from many input formats (Module 2).

Keeps headings, tables, lists, and code blocks where the source format allows.
Extraction failures raise ExtractionError with a clear message.
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from ..logging_config import get_logger

logger = get_logger(__name__)


class ExtractionError(Exception):
    pass


# Default OCR languages: Persian + English (the primary audience). Tesseract
# language packs must be installed for these codes (fas, eng).
DEFAULT_OCR_LANGS = "fas+eng"


def _ocr_ready() -> tuple[bool, str]:
    """Return (available, reason). OCR needs the Python bindings plus the
    Tesseract engine binary; report a precise, actionable reason when not."""
    try:
        import pytesseract  # noqa: F401
        import pypdfium2  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError as e:
        return False, (
            f"OCR support is not installed ({e}). Add pytesseract, pypdfium2 and "
            "Pillow to the backend environment."
        )
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
    except Exception:  # noqa: BLE001 - any failure means the engine is unusable
        return False, (
            "The Tesseract OCR engine is not installed on the server. Install it, "
            "e.g. `apt-get install tesseract-ocr tesseract-ocr-fas tesseract-ocr-eng "
            "poppler-utils`, then retry."
        )
    return True, ""


def _clean_ocr_text(text: str) -> str:
    """Tidy raw OCR output: drop empty lines and trailing whitespace so the
    stored text reads cleanly instead of arriving garbled."""
    lines = [line.strip() for line in (text or "").splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _ocr_image_obj(image, langs: str) -> str:
    import pytesseract

    return _clean_ocr_text(pytesseract.image_to_string(image, lang=langs))


def _from_image(path: Path, langs: str = DEFAULT_OCR_LANGS) -> str:
    ready, reason = _ocr_ready()
    if not ready:
        raise ExtractionError(reason)
    from PIL import Image, ImageOps, UnidentifiedImageError

    try:
        with Image.open(str(path)) as image:
            # EXIF-orient and grayscale for steadier recognition.
            prepared = ImageOps.grayscale(ImageOps.exif_transpose(image))
            text = _ocr_image_obj(prepared, langs)
    except UnidentifiedImageError as e:
        raise ExtractionError(f"Could not read image {path.name}: {e}") from e
    if not text:
        raise ExtractionError(
            "OCR found no readable text in the image. Check the scan quality and "
            "that the correct Tesseract language packs are installed."
        )
    return text


def _ocr_pdf(path: Path, langs: str = DEFAULT_OCR_LANGS) -> str:
    """Rasterize each PDF page and OCR it — for scanned PDFs with no text layer."""
    ready, reason = _ocr_ready()
    if not ready:
        raise ExtractionError(reason)
    import pypdfium2 as pdfium

    parts: list[str] = []
    pdf = pdfium.PdfDocument(str(path))
    try:
        for i in range(len(pdf)):
            page = pdf[i]
            # ~180 DPI: a good accuracy/speed trade-off for OCR.
            bitmap = page.render(scale=2.5)
            image = bitmap.to_pil()
            text = _ocr_image_obj(image, langs)
            if text:
                parts.append(f"\n\n## Page {i + 1}\n\n{text}")
    finally:
        pdf.close()
    if not parts:
        raise ExtractionError("OCR found no readable text in the scanned PDF.")
    return "".join(parts).strip()


def _from_pdf(path: Path, langs: str = DEFAULT_OCR_LANGS) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise ExtractionError("pypdf not installed") from e
    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"\n\n## Page {i}\n\n{text.strip()}")
    combined = "".join(parts).strip()
    # A scanned PDF has no (or a negligible) text layer — fall back to OCR so
    # image-only Persian/English PDFs still produce clean, usable text.
    if len(combined) < 40:
        try:
            return _ocr_pdf(path, langs)
        except ExtractionError:
            if combined:
                return combined  # keep whatever sparse text we did find
            raise
    return combined


def _from_docx(path: Path) -> str:
    from docx import Document as Docx  # python-docx

    doc = Docx(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        style = (para.style.name or "").lower()
        text = para.text.strip()
        if not text:
            continue
        if style.startswith("heading"):
            level = "".join(c for c in style if c.isdigit()) or "2"
            lines.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            lines.append(text)
    for table in doc.tables:
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        if rows:
            lines.append(_table_to_md(rows))
    if not lines:
        raise ExtractionError("DOCX contained no readable text.")
    return "\n\n".join(lines)


def _from_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(cells)
        if rows:
            blocks.append(f"## Sheet: {ws.title}\n\n{_table_to_md(rows)}")
    wb.close()
    if not blocks:
        raise ExtractionError("Workbook had no data.")
    return "\n\n".join(blocks)


def _from_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            slides.append(f"## Slide {i}\n\n" + "\n\n".join(texts))
    if not slides:
        raise ExtractionError("Presentation had no text.")
    return "\n\n".join(slides)


def _from_html(raw: str) -> str:
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    soup = BeautifulSoup(raw, "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
        tag.decompose()
    body = soup.body or soup
    return markdownify(str(body), heading_style="ATX").strip()


def _from_csv(raw: str) -> str:
    reader = csv.reader(io.StringIO(raw))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        raise ExtractionError("CSV had no rows.")
    return _table_to_md(rows)


def _from_json(raw: str) -> str:
    try:
        obj = json.loads(raw)
    except json.JSONDecodeError as e:
        raise ExtractionError(f"Invalid JSON: {e}") from e
    return "```json\n" + json.dumps(obj, indent=2, ensure_ascii=False) + "\n```"


def _table_to_md(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [r + [""] * (width - len(r)) for r in rows]
    header = norm[0]
    md = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for r in norm[1:]:
        md.append("| " + " | ".join(cell.replace("|", "\\|") for cell in r) + " |")
    return "\n".join(md)


EXT_MAP = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".pptx": "pptx",
    ".html": "html",
    ".htm": "html",
    ".csv": "csv",
    ".json": "json",
    ".md": "text",
    ".markdown": "text",
    ".txt": "text",
    ".log": "text",
    # Images go through OCR.
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".tif": "image",
    ".tiff": "image",
    ".bmp": "image",
    ".webp": "image",
    ".gif": "image",
}


def extract_file(path: Path, ocr_langs: str = DEFAULT_OCR_LANGS) -> str:
    """Extract markdown-ish text from a file on disk. Scanned PDFs and image
    files are read with OCR (Tesseract) using `ocr_langs` (default fas+eng)."""
    ext = path.suffix.lower()
    kind = EXT_MAP.get(ext)
    if kind is None:
        # Best effort: try as UTF-8 text.
        kind = "text"
    try:
        if kind == "pdf":
            return _from_pdf(path, ocr_langs)
        if kind == "image":
            return _from_image(path, ocr_langs)
        if kind == "docx":
            return _from_docx(path)
        if kind == "xlsx":
            return _from_xlsx(path)
        if kind == "pptx":
            return _from_pptx(path)
        raw = path.read_text(encoding="utf-8", errors="replace")
        if kind == "html":
            return _from_html(raw)
        if kind == "csv":
            return _from_csv(raw)
        if kind == "json":
            return _from_json(raw)
        return raw.strip()
    except ExtractionError:
        raise
    except Exception as e:  # noqa: BLE001
        logger.exception("Extraction failed for %s", path.name)
        raise ExtractionError(f"Could not extract {path.name}: {e}") from e


def extract_text(raw: str, kind: str = "text") -> str:
    """Extract from an in-memory string (kind: text/html/csv/json/markdown)."""
    kind = kind.lower()
    if kind == "html":
        return _from_html(raw)
    if kind == "csv":
        return _from_csv(raw)
    if kind == "json":
        return _from_json(raw)
    return raw.strip()
